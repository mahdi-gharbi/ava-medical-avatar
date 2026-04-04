# upload_data.py
import streamlit as st
import json
import sys
import os
import io
import chromadb
import pandas as pd
import pdfplumber
from docx import Document
from pptx import Presentation
from chromadb.utils import embedding_functions

# ==============================
# 🔥 FIX IMPORT BACKEND
# ==============================
BASE_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "../../"))
sys.path.append(os.path.join(BASE_DIR, "ai_backend"))

from rag.ingest.ingest_service import ingest_uploaded_data

# ==============================
# 🔥 CHROMA PATH UNIQUE
# ==============================
CHROMA_PATH = os.path.join(BASE_DIR, "ai_backend", "rag", "chroma_db")

# ==============================
# 🔥 EMBEDDING UNIQUE
# ==============================
EMBED_MODEL = "paraphrase-multilingual-MiniLM-L12-v2"

emb_fn = embedding_functions.SentenceTransformerEmbeddingFunction(
    model_name=EMBED_MODEL
)

# ==============================
# 📂 PARSER MULTI-FORMAT
# ==============================
def parse_uploaded_file(uploaded_file) -> list[dict]:
    """
    Parse n'importe quel fichier uploadé → retourne une liste de dicts
    compatibles avec ingest_uploaded_data()
    """
    name = uploaded_file.name
    ext = name.split(".")[-1].lower()

    # ── JSON ──────────────────────────────────────────────────────────────────
    if ext == "json":
        data = json.load(uploaded_file)
        if isinstance(data, dict):
            data = [data]
        return data

    # ── CSV ───────────────────────────────────────────────────────────────────
    elif ext == "csv":
        df = pd.read_csv(uploaded_file)
        return df.fillna("").to_dict(orient="records")

    # ── XLSX ──────────────────────────────────────────────────────────────────
    elif ext == "xlsx":
        df = pd.read_excel(uploaded_file)
        return df.fillna("").to_dict(orient="records")

    # ── TXT ───────────────────────────────────────────────────────────────────
    elif ext == "txt":
        content = uploaded_file.read().decode("utf-8", errors="ignore")
        return [{"name": name, "description": content}]

    # ── PDF ───────────────────────────────────────────────────────────────────
    elif ext == "pdf":
        text = ""
        with pdfplumber.open(uploaded_file) as pdf:
            for page in pdf.pages:
                text += page.extract_text() or ""
        return [{"name": name, "description": text}]

    # ── DOCX ──────────────────────────────────────────────────────────────────
    elif ext == "docx":
        doc = Document(uploaded_file)
        text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
        return [{"name": name, "description": text}]

    # ── PPTX ──────────────────────────────────────────────────────────────────
    elif ext == "pptx":
        prs = Presentation(uploaded_file)
        slides_text = []
        for i, slide in enumerate(prs.slides):
            slide_text = ""
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_text += shape.text + "\n"
            if slide_text:
                slides_text.append(f"Slide {i+1}:\n{slide_text}")
        text = "\n".join(slides_text)
        return [{"name": name, "description": text}]

    else:
        raise ValueError(f"Format non supporté : {ext}")


# ==============================
# 🎨 UI CONFIG
# ==============================
st.set_page_config(page_title="RAG Upload + Viewer", layout="wide")

st.title("📤 Upload + Visualisation RAG")

# ==============================
# 📤 UPLOAD
# ==============================
st.header("📤 Upload JSON")

uploaded_file = st.file_uploader(
    "Upload fichier (JSON, CSV, PDF, DOCX, PPTX, TXT, XLSX)",
    type=["json", "csv", "pdf", "docx", "pptx", "txt", "xlsx"]
)

if uploaded_file:
    try:
        data = parse_uploaded_file(uploaded_file)
        nb_added, total = ingest_uploaded_data(data)

        if nb_added == 0:
            st.warning(f"⚠️ Tous les documents existent déjà ({total} au total)")
        else:
            st.success(f"✅ {nb_added} nouveaux documents ajoutés")
            st.info(f"📊 Total dans global_docs : {total}")

    except Exception as e:
        st.error(f"❌ Erreur : {e}")

# ==============================
# 🔗 CONNECT DB
# ==============================
st.header("📊 Données stockées")

client = chromadb.PersistentClient(path=CHROMA_PATH)

collection_names = [c.name for c in client.list_collections()]
st.write("📁 Collections:", collection_names)

# ==============================
# 📚 VIEW ALL COLLECTIONS
# ==============================
st.header("📚 Toutes les collections + documents")

if not collection_names:
    st.warning("❌ Aucune collection trouvée")
else:
    total_documents = 0

    for name in collection_names:
        st.subheader(f"📂 Collection : {name}")

        # 🔥 get_collection SANS emb_fn pour l'affichage (.get() n'embed rien)
        col = client.get_collection(name=name)

        all_docs = col.get()

        if not all_docs or len(all_docs["ids"]) == 0:
            st.warning("⚠️ Aucun document dans cette collection")
            continue

        doc_count = len(all_docs["ids"])
        total_documents += doc_count

        st.success(f"{doc_count} document(s)")

        # 🔥 FIX SLIDER : évite min == max quand doc_count == 1
        if doc_count == 1:
            max_docs = 1
        else:
            max_docs = st.slider(
                f"Docs à afficher pour {name}",
                1, doc_count, min(20, doc_count),
                key=f"slider_{name}"
            )

        for i in range(max_docs):
            with st.expander(f"📄 Doc {i+1}"):
                st.write("🆔 ID :", all_docs["ids"][i])
                st.write("📦 Metadata :", all_docs["metadatas"][i])
                st.text(all_docs["documents"][i])

    st.success(f"✅ TOTAL GLOBAL : {total_documents} documents")

# ==============================
# 🔍 SEARCH RAG
# ==============================
st.header("🔍 Recherche dans RAG")

search_collection = st.selectbox(
    "Collection à rechercher",
    options=collection_names if collection_names else ["global_docs"],
    index=collection_names.index("global_docs") if "global_docs" in collection_names else 0
)

query = st.text_input("Tape ta recherche")

if query:
    try:
        search_col = client.get_collection(name=search_collection)

        # 🔥 On génère l'embedding manuellement pour éviter le conflit
        query_embedding = emb_fn([query])

        n_results = min(5, search_col.count())

        if n_results == 0:
            st.warning("⚠️ Collection vide")
        else:
            results = search_col.query(
                query_embeddings=query_embedding,
                n_results=n_results
            )

            if not results["documents"][0]:
                st.warning("Aucun résultat trouvé")
            else:
                for j, (doc, meta) in enumerate(
                    zip(results["documents"][0], results["metadatas"][0])
                ):
                    with st.expander(f"🔎 Résultat {j+1}"):
                        st.write("📦 Metadata :", meta)
                        st.text(doc)

    except Exception as e:
        st.error(f"❌ Erreur recherche : {e}")